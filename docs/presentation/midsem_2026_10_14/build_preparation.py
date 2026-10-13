"""Build presentation documents/plots from committed aggregates; no fitting.
Run from repository root: python docs/presentation/midsem_2026_10_14/build_preparation.py
Requires matplotlib, python-docx and reportlab in the document-build environment.
Does not change research dependencies, source, configurations or results.
"""
from pathlib import Path
import csv, hashlib, json, os, re
os.environ.setdefault('MPLCONFIGDIR', '/tmp/midsem-matplotlib' if os.name != 'nt' else str(Path(os.environ['TEMP']) / 'midsem-matplotlib'))
os.environ.setdefault('XDG_CACHE_HOME', os.environ.get('TEMP', '/tmp'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from matplotlib import font_manager
pdfmetrics.registerFont(TTFont('Presentation', font_manager.findfont('DejaVu Sans')))
pdfmetrics.registerFont(TTFont('PresentationBold', font_manager.findfont(font_manager.FontProperties(family='DejaVu Sans', weight='bold'))))
from xml.sax.saxutils import escape

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
PLOTS=OUT/'plots'; PLOTS.mkdir(exist_ok=True)
def read(path):
 with (ROOT/path).open(newline='') as f:return list(csv.DictReader(f))
scores=read('docs/audit/recomputed_scores.csv')
abl=read('docs/audit/recomputed_ablation.csv')
pseudo=read('docs/audit/pseudo_label_audit.csv')
history=read('docs/EC_EBO_SMALL_STUDY_SCORES.csv')
assert len(scores)==120 and len(abl)==20 and len(pseudo)==4 and len(history)==20
methods=['nores','smote','ugo_smote','bfch_smote','ec_ebo','sshr']
names=['NORES','SMOTE','UGO-SMOTE','BFCH-SMOTE','EC-EBO','SSHR']
metrics=['f1','gmean','mcc','probability_auc','fpr']
labels=['F1','G-mean','MCC','Probability ROC-AUC','False-positive rate']
conditions=[(d,c) for d in ['ecoli4','yeast4'] for c in ['LogReg','GNB']]
plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
lookup={(r['dataset'],r['classifier'],r['method'],r['metric']):r for r in scores}
for metric,label in zip(metrics,labels):
 fig,axs=plt.subplots(2,2,figsize=(11,6.7))
 for ax,(d,c) in zip(axs.flat,conditions):
  rows=[lookup[d,c,m,metric] for m in methods]
  ax.errorbar([float(r['mean']) for r in rows],range(6),xerr=[float(r['sample_sd']) for r in rows],fmt='o',capsize=3,color='#176b87')
  ax.set_yticks(range(6),names);ax.invert_yaxis();ax.set_title(f'{d.title()} / {c}');ax.grid(axis='x',alpha=.2)
  ax.set_xlabel(label+(' (lower is better)' if metric=='fpr' else ' (higher is better)'))
 fig.suptitle(f'Six-method pilot: {label}',fontsize=16)
 fig.text(.5,.015,'Mean ± sample SD over 25 fold events; dependent folds, not confidence intervals. SD bars may extend outside metric bounds.',ha='center',fontsize=8)
 fig.tight_layout(rect=(0,.04,1,.94));fig.savefig(PLOTS/f'{metric}.png',dpi=170);plt.close(fig)
fig,axs=plt.subplots(1,5,figsize=(15,4.5))
for ax,metric,label in zip(axs,metrics,labels):
 rows=[next(r for r in abl if (r['dataset'],r['classifier'],r['metric'])==(d,c,metric)) for d,c in conditions]
 ax.barh(range(4),[float(r['delta']) for r in rows],color=['#c66a38' if float(r['delta'])<0 else '#176b87' for r in rows])
 ax.set_yticks(range(4),[f'{d.title()}\n{c}' for d,c in conditions]);ax.invert_yaxis();ax.axvline(0,color='black',lw=.8);ax.set_title(label,fontsize=10)
 ax.tick_params(axis='x',labelsize=8)
fig.suptitle('SSHR paired effect: with U minus empty U',fontsize=16)
fig.text(.5,.015,'Mean paired differences, 25 pairs per condition. F1 fell and FPR rose in all four; some other metrics improved. No significance claim.',ha='center',fontsize=9)
fig.tight_layout(rect=(0,.05,1,.91));fig.savefig(PLOTS/'sshr_ablation.png',dpi=170);plt.close(fig)
fig,ax=plt.subplots(figsize=(8,4))
pr=[next(r for r in pseudo if r['dataset']==d and r['classifier']=='LogReg') for d in ['ecoli4','yeast4']]
vals=[100*float(r['precision']) for r in pr]
ax.bar(['Ecoli4','Yeast4'],vals,color=['#176b87','#c66a38']);ax.set_ylim(0,100);ax.set_ylabel('Correct minority pseudo-labels (%)')
for i,r in enumerate(pr):ax.text(i,vals[i]+3,f"{r['correct']}/{r['selected']} = {vals[i]:.1f}%",ha='center')
ax.set_title('SSHR selected-row correctness')
fig.text(.5,.02,'Selection events across folds; same counts for each classifier. Not unique rows or synthetic-vector accuracy.',ha='center',fontsize=8)
fig.tight_layout(rect=(0,.07,1,1));fig.savefig(PLOTS/'pseudo_label_precision.png',dpi=170);plt.close(fig)
# Separate historical study: never pool its ten folds with the reduced-L pilot.
fig,axs=plt.subplots(2,2,figsize=(11,6.5))
for ax,(d,c) in zip(axs.flat,[(d,c) for d in ['winequality_red_8_vs_6','vowel0_acoustic10_grouped'] for c in ['LogReg','GNB']]):
 rows=[next(r for r in history if (r['dataset_id'],r['classifier'],r['method'])==(d,c,m)) for m in methods[:-1]]
 ax.errorbar([float(r['f1_mean']) for r in rows],range(5),xerr=[float(r['f1_sample_sd']) for r in rows],fmt='o',capsize=3,color='#176b87')
 ax.set_yticks(range(5),names[:-1]);ax.invert_yaxis();ax.set_title(f'{"Winequality" if d.startswith("wine") else "Vowel acoustic"} / {c}');ax.set_xlabel('F1')
fig.suptitle('Historical five-method study: separate protocol',fontsize=16)
fig.text(.5,.015,'Mean ± sample SD over ten folds; Vowel speaker groups. Do not pool with Ecoli4/Yeast4 or treat as independent replication of BFCH study.',ha='center',fontsize=8)
fig.tight_layout(rect=(0,.05,1,.94));fig.savefig(PLOTS/'historical_f1.png',dpi=170);plt.close(fig)

intro='''# Mid-Sem Evaluation Preparation — 14 October 2026

Project: Imbalanced Learning Reproductions (UGO, BFCH-SMOTE, EC-EBO and SSHR).
Repository: MAYANK-RATHVA05/rb_ugo. Working branch: sshr-implementation.
Prepared from saved study aggregates; no new training was run for this presentation.

## 1. What to say at the start

“Our project investigates class imbalance through audited implementations of four published resampling methods. We compare them under common evaluation conditions and study whether unlabelled training observations help SSHR. Our current contribution is reproduction engineering and controlled empirical analysis; a novel oversampling algorithm has not been established.”

Evidence: docs/PROJECT_SCOPE.md; docs/audit/claims_table.md; docs/audit/audit_protocol.md.

## 2. Three research questions and progress

| Question | Implemented work | What remains unproved |
| --- | --- | --- |
| A: Can boundary-focused generation avoid unnecessary overlap? | UGO and BFCH-SMOTE; training-only distance proxy | No guarantee of true overlap avoidance or correct synthetic labels |
| B: Can intra-class imbalance be handled? | EC-EBO clustering and sample allocation | Held-out minority-subregion improvement was not measured |
| C: Can unlabelled observations support minority augmentation? | SSHR clustering, pseudo-label selection and SMOTE fallback; paired U ablation | Broad benefit and exact author-code equivalence remain unverified |

SSHR admits existing unlabelled observations with inferred minority labels. Only its SMOTE fallback generates new feature vectors. Source: docs/SSHR_METHOD_DECISIONS.md, Algorithm 2 audit; docs/audit/audit_protocol.md.

## 3. Project Review suggestions: honest status

The official Project Review comments were not supplied in the available records. Before presenting, insert their exact wording and confirm the mapping below. Do not describe a suggestion as completed merely because it appears in this preparation guide.

| Item from available project discussions | Current evidence/status | Action before evaluation |
| --- | --- | --- |
| Reproduce previous papers before claiming an extension | Four audited implementations; all classified APPROXIMATION | Explain declared source ambiguities and departures |
| Compare methods fairly | Common L/folds/classifiers and training-only preprocessing | Show protocol and audit evidence |
| Investigate unlabelled data | Paired U/empty-U study and correctness audit complete | Show ablation and precision plots |
| Local safety and inverse-density extension | Proposal only; novelty unverified | Do not present as implemented or a measured result |
| Official reviewer-specific suggestion | UNVERIFIED: enter exact comment | Record completed/partial/pending plus a file or experiment as evidence |

Source: docs/audit/audit_protocol.md and contribution_and_qa.md; project discussion. UGO already relocates to the nearest clean same-class point, not a centroid (src/noise_detection.py; UGO Section 3.4). A centroid replacement is therefore not an established new contribution.

## 4. Explain the main experiment in one minute

Two numeric datasets: Ecoli4 and Yeast4. Two final classifiers: Logistic Regression (LogReg) and Gaussian Naive Bayes (GNB). Six methods: NORES, SMOTE, UGO-SMOTE, BFCH-SMOTE, EC-EBO and SSHR.

Five repeats of five outer folds give 25 held-out fold events per condition. Approximately half the training labels are retained as observed L; training-only U contains the other rows. Every baseline gets the same reduced L. Raw identical-coordinate groups stay together across train/test and L/U. Preprocessing fits L only; hidden labels audit selections afterward. Test data does not enter generation or parameter selection in the inspected pipeline.

The main matrix has 600 successful jobs. Empty-U SSHR and duplicated NORES controls add 200 jobs. All 800 saved jobs were audited. These are repeated fits, not 800 independent datasets. Sources: scripts/run_sshr_small_study.py; docs/audit/audit_protocol.md; docs/audit/audit_verification.json.

Metric meanings: F1 balances minority precision/recall; G-mean balances sensitivity/specificity; MCC summarizes confusion-matrix association; probability ROC-AUC measures ranking; FPR measures majority false alarms. Higher is better except FPR. Error bars below are sample SD, not confidence intervals or significance tests.

## 5. Main result: no universal winner

UGO-SMOTE has the highest mean F1 on Ecoli4 for both classifiers; EC-EBO has the highest mean F1 on Yeast4 for both classifiers. This is a descriptive ranking in this pilot. Sources: docs/audit/recomputed_scores.csv; docs/audit/audit_numbers.md.

![Main F1 comparison](plots/f1.png)

For LogReg, Ecoli4 F1 is 0.8890 for UGO-SMOTE versus 0.8720 for NORES. Yeast4 F1 is 0.3189 for EC-EBO versus 0.2018 for NORES. Yeast4 LogReg FPR also rises from 0.0046 to 0.0923: report the cost alongside the gain.

![False-positive rates](plots/fpr.png)

## 6. Question C: what changed with unlabelled data?

With U minus empty U, mean SSHR F1 fell in all four conditions: Ecoli4/LogReg -0.0777, Ecoli4/GNB -0.0922, Yeast4/LogReg -0.0096, Yeast4/GNB -0.0054. FPR rose in all four. Yeast4/GNB ROC-AUC and G-mean improved, so “unlabelled data never helps” is unsupported. Source: docs/audit/recomputed_ablation.csv.

![Paired SSHR ablation](plots/sshr_ablation.png)

![Pseudo-label correctness](plots/pseudo_label_precision.png)

Correctness is 170/313 (54.3%) for Ecoli4 and 96/273 (35.2%) for Yeast4, separately for each classifier. Counts are selection events over repeated folds. Wrong pseudo-labels are a plausible contributor to harm, but this joint U ablation does not isolate the cause: U also changes clustering, cleaning and fallback. Source: docs/audit/pseudo_label_audit.csv; docs/audit/contribution_and_qa.md.

## 7. Contribution, limitations and next step

Completed contributions: auditable implementations with explicit paper ambiguities; a common controlled evaluation; a paired local U ablation; and retrospective quality/count/shortfall diagnostics. No novel method is established. Do not call a local observation first-in-literature without a prior-art review.

All four whole-method implementations are classified APPROXIMATION. SSHR uses a literal printed objective that can exceed one; held-out standard F1 is a different quantity and was independently verified. Two datasets and two classifiers cannot establish general superiority. Repeated folds share observations. No between-method significance analysis or tuned-threshold NORES control is available. BFCH has a distance proxy, not a true-overlap guarantee; EC-EBO lacks held-out subgroup recall. Sources: docs/audit/audit_protocol.md; claims_table.md.

Proposed next step — HYPOTHESIS ONLY: test whether an uncertainty-based admission filter improves SSHR pseudo-label quality and FPR without unacceptable loss of recall or candidate coverage. Resolve source ambiguities, check closest prior art and predeclare an independent confirmation protocol first. Do not tune against these already-inspected test outcomes. The broader local-safety/inverse-density idea remains unimplemented and its novelty is unverified.

## 8. Five-minute speaking plan

| Time | Show | What to say |
| --- | --- | --- |
| 0:00–0:35 | Questions A/B/C | State the three questions; do not claim they are solved |
| 0:35–1:05 | Scope and implementation files | Four published methods plus two controls; declared approximations |
| 1:05–1:55 | Protocol | Same L/folds; training-only U; 600+200 audited jobs |
| 1:55–2:45 | F1 and FPR plots | Mean leaders vary; report false-positive cost |
| 2:45–3:30 | Ablation and precision plots | F1 falls with U; low purity; other metrics can improve |
| 3:30–4:05 | Limitations | No significance, no novelty established, source ambiguities |
| 4:05–4:35 | Contribution | Reproduction, fair comparison and diagnostic analysis |
| 4:35–5:00 | Next step | A falsifiable hypothesis, not a completed result |

## 9. Questions to rehearse

| Supervisor question | Honest answer |
| --- | --- |
| What is new? | A project-level controlled analysis; algorithmic novelty is not established |
| Why does U hurt? | Incorrect pseudo-labels are plausible, but the experiment changes several components and does not prove the cause |
| Is the difference significant? | We report means and descriptive SDs; no between-method significance claim |
| Did you reproduce the complete papers? | No: declared implementation choices and small pilots; the complete UGO benchmark is unfinished |
| Why only two datasets here? | A controlled diagnostic pilot; broader confirmation remains necessary |
| Is SSHR generating new samples? | It mainly selects existing U rows; only fallback SMOTE generates new vectors |
| Have A and B been solved? | The methods target them; true-overlap guarantees and subgroup benefit are not established |
| Are hidden/test labels used? | Hidden labels design the simulation and later audit correctness, not sampler learning; test labels score final predictions |
| Where is the evidence? | CSV aggregates, method-decision documents and the independent audit |
| What did you personally do? | Describe each member’s actual work and credit original authors and tools; individual authorship was not audited |

More detailed answers: docs/audit/contribution_and_qa.md.

## 10. Antigravity demonstration and backup checklist

Use the local checkout C:\\ugo\\ugo_base and branch sshr-implementation. Cloud changes do not automatically update this folder. In its terminal run git status, git fetch origin, then git switch sshr-implementation and git pull --ff-only origin sshr-implementation if the working tree permits. If the branch does not exist locally, use git switch --track origin/sshr-implementation. If Git reports local changes or divergence, stop and preserve them; do not reset or force-pull.

Open this folder first: docs/presentation/midsem_2026_10_14/. Open Mid_Sem_Preparation.pdf in the system PDF viewer and keep the plots/ folder ready. In the editor, open README.md, docs/SSHR_SMALL_STUDY.md and docs/audit/audit_numbers.md; use Markdown preview if available. Open docs/sshr_pilot/ecoli4/scores.csv or yeast4/scores.csv when asked for underlying tables. No Python environment or retraining is needed to view these committed outputs.

For implementation questions open src/ugo.py, src/bfch_smote.py, src/ec_ebo.py, src/sshr.py and src/sshr_clustering.py. For run configuration open scripts/run_sshr_small_study.py. For diagnostics open each dataset’s diagnostics.csv, clusters.csv and ablation.csv under docs/sshr_pilot/.

Before leaving: confirm branch and latest commit; open the PDF and every figure offline; check official review suggestions; assign speaking roles honestly; copy the PDF/plots to a backup drive. Keep raw data, individual predictions and local paths out of projected screenshots. Do not launch the full benchmark during evaluation.

## 11. Inventory of other completed and partial experiments

These experiments used different protocols. Show them as separate backup evidence; do not merge their averages or count repeated control runs as independent replication.

| Study | Recorded coverage | Where to show it |
| --- | --- | --- |
| UGO simulation pilot | 300 successful jobs; reconstructed linear/nonlinear inputs; three classifiers | docs/reproduction/SIMULATION_PILOT.md |
| BFCH comparison | Winequality/Vowel; four methods; two classifiers; ten folds; 160 successful jobs | docs/BFCH_SMALL_STUDY.md and BFCH_SMALL_STUDY_SCORES.csv |
| EC-EBO comparison | Same two datasets; five methods; two classifiers; ten folds; 200 successful jobs | docs/EC_EBO_SMALL_STUDY.md and EC_EBO_SMALL_STUDY_SCORES.csv |
| Controlled six-method pilot | Ecoli4/Yeast4; 600 main + 200 ablation/control jobs | docs/SSHR_SMALL_STUDY.md; docs/sshr_pilot/ |
| Historical larger UGO benchmark | Interrupted: 28,919 saved jobs reported; 18 datasets terminal, including failures; full 74 incomplete | docs/reproduction/BENCHMARK_EXECUTION.md; original cache not in Git |

The historical inventory is based on recorded reports; this preparation build does not re-audit those earlier individual runs. A fresh Windows clone will have the tracked reports, not ignored job caches/raw datasets. Historical terminal jobs include failures; do not claim all succeeded.

![Historical F1 comparison](plots/historical_f1.png)

## 12. Complete main-pilot metric table

Mean ± sample SD over 25 fold events. Sources: docs/audit/recomputed_scores.csv and original docs/sshr_pilot/<dataset>/scores.csv. All five metrics are shown; plots for G-mean, MCC and ROC-AUC are also in plots/.
'''
parts=[intro]
for d,c in conditions:
 parts += [f'### {d.title()} / {c}','| Method | F1 | G-mean | MCC | ROC-AUC | FPR |','| --- | --- | --- | --- | --- | --- |']
 for m,name in zip(methods,names):
  vals=[f"{float(lookup[d,c,m,k]['mean']):.4f} ± {float(lookup[d,c,m,k]['sample_sd']):.4f}" for k in metrics]
  parts.append('| '+' | '.join([name]+vals)+' |')
parts+=['## 13. Historical five-method metric table','Separate ten-fold protocol. Mean values below; sample SDs and source hashes are in docs/EC_EBO_SMALL_STUDY_SCORES.csv. The preceding BFCH study retains its own full tables in docs/BFCH_SMALL_STUDY_SCORES.csv.']
for d,c in [(d,c) for d in ['winequality_red_8_vs_6','vowel0_acoustic10_grouped'] for c in ['LogReg','GNB']]:
 parts +=[f'### {d} / {c}','| Method | F1 | G-mean | MCC | AUC | FPR |','| --- | --- | --- | --- | --- | --- |']
 for m,name in zip(methods[:-1],names[:-1]):
  r=next(r for r in history if (r['dataset_id'],r['classifier'],r['method'])==(d,c,m))
  parts.append('| '+' | '.join([name]+[f'{float(r[k+"_mean"]):.4f}' for k in ['f1','gmean','mcc','auc','fpr']])+' |')
parts+=['## 14. UGO simulation results (historical report)','Mean MCC from docs/reproduction/SIMULATION_PILOT.md; reconstructed data, not author coordinates. Other simulation diagnostics and resource limitations remain in that report.']
sim=(ROOT/'docs/reproduction/SIMULATION_PILOT.md').read_text().splitlines()
parts += [l for l in sim if l.startswith('|')]
parts+=['## 15. Source and build record','All new plots use saved aggregates. source_manifest.json records hashes of the source tables/reports used by this build. Rebuild from the repository root with: python docs/presentation/midsem_2026_10_14/build_preparation.py. Document-only dependencies: matplotlib, python-docx, reportlab. No new experiment or source change is implied.','Detailed paper citations: docs/audit/paper_sources.md. Numerical audit: docs/audit/audit_numbers.md. Protocol/fidelity: docs/audit/audit_protocol.md. Existing study reports remain the authoritative historical records.']
md='\n\n'.join(parts)+'\n'
# Consecutive table lines must remain adjacent in Markdown.
md=re.sub(r'(\|)\n\n(?=\|)',r'\1\n',md)
(OUT/'PREPARATION.md').write_text(md)
# Shared simple Markdown renderer for Word and PDF.
doc=Document();doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet()
for style in styles.byName.values():
 style.fontName='PresentationBold' if style.name.startswith(('Heading','Title')) else 'Presentation'
styles['BodyText'].fontSize=8.5;styles['BodyText'].leading=12
story=[];lines=md.splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('|'):
  rows=[]
  while True:
   if not re.fullmatch(r'[| :\-]+',line):rows.append([v.strip() for v in line.strip('|').split('|')])
   if i>=len(lines) or not lines[i].strip().startswith('|'):break
   line=lines[i].strip();i+=1
  t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Light Shading Accent 1'
  for j,v in enumerate(rows[0]):t.rows[0].cells[j].text=v
  for row in rows[1:]:
   for cell,v in zip(t.add_row().cells,row):cell.text=v
  for row in t.rows:
   for cell in row.cells:
    for par in cell.paragraphs:
     for run in par.runs:run.font.size=Pt(8)
  rt=Table([[Paragraph(escape(v),styles['BodyText']) for v in row] for row in rows],repeatRows=1,colWidths=[(A4[0]-72)/len(rows[0])]*len(rows[0]))
  rt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
  story +=[rt,Spacer(1,10)];continue
 if line.startswith('!['):
  path=OUT/re.search(r'\]\(([^)]+)\)',line).group(1)
  doc.add_picture(str(path),width=Inches(6.2))
  from PIL import Image as PILImage
  with PILImage.open(path) as im:w,h=im.size
  story +=[Image(str(path),width=A4[0]-72,height=(A4[0]-72)*h/w),Spacer(1,10)];continue
 if line.startswith('#'):
  level=len(line)-len(line.lstrip('#'));title=line[level:].strip()
  doc.add_heading(title,level=min(level,3))
  story.append(Paragraph(escape(title),styles['Title' if level==1 else 'Heading'+str(min(level-1,3))]));continue
 doc.add_paragraph(line);story +=[Paragraph(escape(line),styles['BodyText']),Spacer(1,7)]
doc.save(OUT/'Mid_Sem_Preparation.docx')
def footer(canvas,document):
 canvas.setFont('Presentation',8);canvas.drawString(36,20,'Mid-Sem preparation | 14 October 2026 | Descriptive research results');canvas.drawRightString(A4[0]-36,20,str(document.page))
SimpleDocTemplate(str(OUT/'Mid_Sem_Preparation.pdf'),title='Mid-Sem Evaluation Preparation - 14 October 2026',pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36).build(story,onFirstPage=footer,onLaterPages=footer)
sources=['docs/audit/recomputed_scores.csv','docs/audit/recomputed_ablation.csv','docs/audit/pseudo_label_audit.csv','docs/EC_EBO_SMALL_STUDY_SCORES.csv','docs/BFCH_SMALL_STUDY_SCORES.csv','docs/reproduction/SIMULATION_PILOT.md','docs/reproduction/BENCHMARK_EXECUTION.md','docs/audit/audit_protocol.md','docs/SSHR_SMALL_STUDY.md']
(OUT/'source_manifest.json').write_text(json.dumps({'purpose':'Presentation derived from existing results; no fitting','sources':{s:hashlib.sha256((ROOT/s).read_bytes()).hexdigest() for s in sources},'plot_count':len(list(PLOTS.glob('*.png')))},indent=2)+'\n')
print('Built Markdown, Word, PDF and 8 plots from saved results.')
